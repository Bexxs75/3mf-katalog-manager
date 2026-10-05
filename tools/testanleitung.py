#!/usr/bin/env python3
"""Test guide generator for preview builds.

Builds a test guide in German or English (HTML, optionally a PDF via headless
Chromium) and a formatted result sheet (.xlsx) per platform, from fixed
building blocks, the "Unreleased" part of CHANGELOG.md in that language and a
per-version scenario file. Run with `-s tools` so `import release_assets`
resolves. Needs openpyxl for the result sheet.
"""
import base64, copy, dataclasses, html, json, pathlib, re, shutil, subprocess, sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import release_assets

PRODUCT = "3MF Katalog Manager Preview"
IDENTIFIER = "com.thebexxs.mfkatalogmanager.preview"
# Tests that need more than clicking around (broken files, PowerShell, log details).
EXPERT_TESTS = {"E5", "K1", "K2", "K3", "K4", "K5", "K6", "P1", "F3", "F4", "F5", "F6", "F7", "N3", "N5"}
# Tests that only apply to some testers (own hardware, extra test data) belong to the full test only.
FULL_ONLY_PREFIXES = ("PB", "GK")
NEW_CHAPTER_RE = re.compile(r"(?:Neu in|New in) (\d+(?:\.\d+)*(?:-\d+)?)")
REPO = "Bexxs75/3mf-katalog-manager"
# Test results go to this alias, where they are triaged automatically.
TEST_EMAIL = "testing@3mfkatalog.de"
PLATFORM_NAMES = {"windows": "Windows", "macos": "macOS", "linux": "Linux"}
# (release_assets platform key, package extension, data-folder path template,
# log-folder path template). Data-folder paths mirror
# docs/benutzerhandbuch/BENUTZERHANDBUCH.md, with the preview identifier. The
# log folder is a separate location: it comes from Tauri's app_log_dir (see
# the tauri_plugin_log::TargetKind::LogDir target in src-tauri/src/lib.rs),
# which on Windows and macOS is not simply "<data-folder>/logs".
PLATFORM_INFO = {
    # Spelled out instead of %APPDATA%: testers read the variable as part of the path.
    "windows": ("Windows-x64", ".msi", r"C:\Users\<dein Name>\AppData\Roaming\{id}", r"C:\Users\<dein Name>\AppData\Local\{id}\logs"),
    "macos": ("macOS-universal", ".dmg", "~/Library/Application Support/{id}", "~/Library/Logs/{id}"),
    "linux": ("Linux-x86_64", ".AppImage", "~/.local/share/{id}", "~/.local/share/{id}/logs"),
}


# Everything that differs between the German and the English guide and sheet.
# Scenario files are `docs/tests/<version>.md` (German) and
# `docs/tests/<version>.en.md`; building blocks live in `bausteine/` and
# `bausteine/en/`.
LANGS = {
    "de": {
        "scenario_suffix": ".md",
        "bausteine": "bausteine",
        "bilder": "bilder",
        "result_image": "So sieht es richtig aus",
        "guide_file": "Testanleitung-{version}-{platform}",
        "sheet_file": "Ergebnisbogen-{version}-{platform}.xlsx",
        "guide_title": "{produkt} {version} – Testanleitung ({platform})",
        "form": "https://3mfkatalog.de/fehler-melden.html",
        "headings": ("Was ist neu", "Installieren", "Datenorte", "Fehler melden", "Deinstallieren"),
        "testdata_heading": "Testdaten vorbereiten",
        "download_one": "Lade `{paket}` von der [Vorschau-Release-Seite]({release}) herunter.",
        "steps": "Schritte:",
        "expected": "Erwartet:",
        "note": "Hinweis:",
        "sheet_intro": ("Trag deine Ergebnisse im Ergebnisbogen ein ({sheet}). Er öffnet sich mit Excel, "
                        "Numbers oder LibreOffice. Diese Anleitung ist nur zum Lesen. Den ausgefüllten Bogen "
                        "schickst du per E-Mail an " + TEST_EMAIL + "."),
        "result_hint": "Ergebnis im Ergebnisbogen eintragen, Zeile {id}",
        "sheet_title": "{produkt} {version} · Ergebnisbogen {platform}",
        "sheet_tab": "Ergebnisbogen",
        "sheet_help": ("Wähle bei jedem Test in der Spalte „Ergebnis“ OK, Fehler oder Übersprungen. Bei „Fehler“ "
                       "bitte kurz notieren, was passiert ist. Die Nummern entsprechen der Testanleitung (PDF)."),
        "columns": ("Nr.", "Test", "Ergebnis", "Notiz"),
        "results": ("OK", "Fehler", "Übersprungen"),
        "invalid_result": "Bitte OK, Fehler oder Übersprungen wählen.",
        "summary": "Zusammenfassung",
        "open": "Noch offen",
        "fields": {
            "windows": [("Gerät", "z. B. Lenovo ThinkPad T14", None),
                        ("Windows-Version", "z. B. Windows 11 24H2", None),
                        ("Virenschutz", "z. B. Microsoft Defender, Avast", None)],
            "macos": [("Gerät", "z. B. MacBook Air M4", None),
                      ("Chip", None, ("Apple Silicon", "Intel")),
                      ("macOS-Version", "z. B. macOS 15.5", None)],
            "linux": [("Gerät", "z. B. Tuxedo InfinityBook", None),
                      ("Distribution und Version", "z. B. Fedora 44, Linux Mint 22", None),
                      ("Desktop", "z. B. KDE Plasma, GNOME, Cinnamon", None),
                      ("Sitzung", None, ("Wayland", "X11"))],
            "common": [("Datum", None, None),
                       ("Name oder Discord (freiwillig)", None, None)],
        },
    },
    "en": {
        "scenario_suffix": ".en.md",
        "bausteine": "bausteine/en",
        "bilder": "bilder/en",
        "result_image": "This is what it should look like",
        "guide_file": "Test-Guide-{version}-{platform}",
        "sheet_file": "Result-Sheet-{version}-{platform}.xlsx",
        "guide_title": "{produkt} {version} – Test guide ({platform})",
        "form": "https://3mfkatalog.de/en/report-a-bug.html",
        "headings": ("What's new", "Installing", "Where the data is", "Reporting a bug", "Uninstalling"),
        "testdata_heading": "Preparing the test data",
        "download_one": "Download `{paket}` from the [preview release page]({release}).",
        "steps": "Steps:",
        "expected": "Expected:",
        "note": "Note:",
        "sheet_intro": ("Enter your results in the result sheet ({sheet}). It opens in Excel, Numbers or "
                        "LibreOffice. This guide is for reading only. Send the filled-in sheet by e-mail to "
                        + TEST_EMAIL + "."),
        "result_hint": "Enter the result in the result sheet, row {id}",
        "sheet_title": "{produkt} {version} · Result sheet {platform}",
        "sheet_tab": "Result sheet",
        "sheet_help": ("For every test, choose OK, Failed or Skipped in the column \"Result\". For \"Failed\", "
                       "please note briefly what happened. The numbers match the test guide (PDF)."),
        "columns": ("No.", "Test", "Result", "Note"),
        "results": ("OK", "Failed", "Skipped"),
        "invalid_result": "Please choose OK, Failed or Skipped.",
        "summary": "Summary",
        "open": "Still open",
        "fields": {
            "windows": [("Device", "e.g. Lenovo ThinkPad T14", None),
                        ("Windows version", "e.g. Windows 11 24H2", None),
                        ("Antivirus", "e.g. Microsoft Defender, Avast", None)],
            "macos": [("Device", "e.g. MacBook Air M4", None),
                      ("Chip", None, ("Apple Silicon", "Intel")),
                      ("macOS version", "e.g. macOS 15.5", None)],
            "linux": [("Device", "e.g. Tuxedo InfinityBook", None),
                      ("Distribution and version", "e.g. Fedora 44, Linux Mint 22", None),
                      ("Desktop", "e.g. KDE Plasma, GNOME, Cinnamon", None),
                      ("Session", None, ("Wayland", "X11"))],
            "common": [("Date", None, None),
                       ("Name or Discord (optional)", None, None)],
        },
    },
}


# Texts of the interactive test assistant (one test per screen, results saved
# as a text file). Kept apart from LANGS because only the assistant uses them.
ASSISTANT_TEXT = {
    "de": {
        "file": "Testassistent-{version}-{platform}.html",
        "result_file": "Testergebnis-{version}-{platform}.txt",
        "title": "{produkt} {version} – Testassistent ({platform})",
        "eyebrow": "Testassistent",
        "hello": "Hallo! Schön, dass du mittestest.",
        "sub": "Wir gehen die Test-Version zusammen durch, Schritt für Schritt.",
        "intro": [["Ein Test nach dem anderen", "Mach die Schritte der Reihe nach und sag danach, ob es geklappt hat."],
                  ["Pause? Kein Problem", "Öffnest du diese Datei wieder, geht es an derselben Stelle weiter."],
                  ["Klappt etwas nicht?", "Schreib kurz dazu, was stattdessen passiert ist. Genau das hilft uns."],
                  ["Zum Schluss", "Ergebnis als Datei speichern und per E-Mail an {email} schicken."]],
        "start": "Los geht’s", "resume": "Weitermachen",
        "prep": "Vorbereitung", "install": "App installieren", "prep_done": "Erledigt, weiter",
        "test_of": "Test {n} von {total}", "all_done": "Alle Tests erledigt",
        "mode_q": "Wie viel Zeit hast du?",
        "modes": {"new": ["Nur das Neue", "Die Tests zu dieser Test-Version und das Update."],
                  "short": ["Kurztest", "Einrichten, Grundfunktionen, Neues und Update."],
                  "full": ["Volltest", "Alles, am Ende optional die Profi-Tests."]},
        "mode_meta": "{n} Tests · ca. {min} Min.",
        "mode_returning": "Du hast schon eine frühere Test-Version getestet, deshalb reicht meist „Nur das Neue“.",
        "mode_update_note": "Dazu kommen jeweils ein paar kurze Update-Tests, sobald die nächste Test-Version erscheint.",
        "mode_label": "Umfang",
        "profi_chapter": "Für Profis (optional)",
        "chapters_btn": "Kapitel", "chapters_title": "Zu einem Kapitel springen",
        "chapters_hint": "Übersprungene Tests bleiben offen. Du kannst jederzeit hierher zurück und sie nachholen.",
        "tip": "Tipp: Klick einen Schritt an, wenn du ihn erledigt hast.",
        "expected": "Das sollte passieren", "note": "Hinweis", "question": "Hat es geklappt?",
        "result_image": "So sieht es richtig aus", "zoom_close": "Klicken oder Esc zum Schließen",
        "download_btn": "Installationsdatei herunterladen", "download_note": "Die Datei landet in deinem Download-Ordner.",
        "time_total": "Aktive Testzeit: {min} Minuten", "time_one": "{min} Min.", "time_lt1": "unter 1 Min.",
        "na": "nicht zutreffend (anderer Testumfang)",
        "ok": "Klappt", "bad": "Klappt nicht", "skip": "Überspringen",
        "why": "Was ist stattdessen passiert?", "why_ph": "z. B. „Es kam die Meldung …“ oder „Der Ordner erschien nicht links“",
        "save_next": "Speichern und weiter", "skip_q": "Warum überspringst du diesen Test?", "choose": "Bitte wählen …",
        "skip_reasons": ["Das kann ich nicht testen (z. B. fehlt mir das Nötige)", "Ich weiß nicht, wie das geht", "Anderer Grund"],
        "next": "Weiter", "back": "← Zurück",
        "done_eyebrow": "Geschafft", "done": "Danke! Alle Tests sind durch.",
        "done_sub": "Du hast uns damit richtig geholfen. Jetzt nur noch das Ergebnis speichern und abschicken.",
        "labels": {"ok": "klappt", "bad": "klappt nicht", "skip": "übersprungen", "open": "offen"},
        "device": "Zum Schluss noch kurz zu deinem Gerät (freiwillig):",
        "final": "Letzter Schritt: Speichere das Ergebnis und schick die Datei per E-Mail an {email}.",
        "mail_btn": "E-Mail-Programm öffnen", "mail_subject": "Testergebnis {version} {platform}",
        "mail_body": "Hallo,\n\nim Anhang mein Testergebnis.\n\n(Bitte die Datei {file} aus dem Download-Ordner anhängen.)",
        "mail_hint": "Die Datei hängst du selbst an, sie liegt in deinem Download-Ordner. Kein E-Mail-Programm? Schick sie einfach von deinem Mail-Konto im Browser an {email}.",
        "save": "Ergebnis speichern", "saved": "Gespeichert als „{file}“ in deinem Download-Ordner.",
        "back_last": "← Letzten Test ändern", "reset": "Von vorn beginnen",
        "reset_confirm": "Wirklich alle Antworten löschen?", "reset_yes": "Ja, alles löschen", "reset_no": "Abbrechen",
        "result_head": "Testergebnis", "status": {"ok": "OK", "bad": "FEHLER", "skip": "ÜBERSPRUNGEN", "open": "OFFEN", "na": "NICHT ZUTREFFEND"},
    },
    "en": {
        "file": "Test-Assistant-{version}-{platform}.html",
        "result_file": "Test-Result-{version}-{platform}.txt",
        "title": "{produkt} {version} – Test assistant ({platform})",
        "eyebrow": "Test assistant",
        "hello": "Hi! Great that you're testing with us.",
        "sub": "We'll go through the test version together, step by step.",
        "intro": [["One test at a time", "Do the steps in order, then tell us whether it worked."],
                  ["Need a break? No problem", "When you open this file again, it continues where you left off."],
                  ["Something doesn't work?", "Write briefly what happened instead. That's exactly what helps us."],
                  ["At the end", "Save the result as a file and send it by e-mail to {email}."]],
        "start": "Let's go", "resume": "Continue",
        "prep": "Preparation", "install": "Install the app", "prep_done": "Done, next",
        "test_of": "Test {n} of {total}", "all_done": "All tests done",
        "mode_q": "How much time do you have?",
        "modes": {"new": ["Just what's new", "The tests for this test version and the update."],
                  "short": ["Short test", "Setup, basics, what's new and the update."],
                  "full": ["Full test", "Everything, with optional expert tests at the end."]},
        "mode_meta": "{n} tests · about {min} min",
        "mode_returning": "You already tested an earlier test version, so \"Just what's new\" is usually enough.",
        "mode_update_note": "Each option also has a few short update tests once the next test version is out.",
        "mode_label": "Scope",
        "profi_chapter": "For experts (optional)",
        "chapters_btn": "Chapters", "chapters_title": "Jump to a chapter",
        "chapters_hint": "Tests you jump over stay open. You can come back here at any time and do them later.",
        "tip": "Tip: click a step once you have done it.",
        "expected": "This should happen", "note": "Note", "question": "Did it work?",
        "result_image": "This is what it should look like", "zoom_close": "Click or press Esc to close",
        "download_btn": "Download the installer", "download_note": "The file goes to your Downloads folder.",
        "time_total": "Active test time: {min} minutes", "time_one": "{min} min", "time_lt1": "under 1 min",
        "na": "not applicable (other test scope)",
        "ok": "Works", "bad": "Doesn't work", "skip": "Skip",
        "why": "What happened instead?", "why_ph": "e.g. \"The message … appeared\" or \"The folder didn't show up on the left\"",
        "save_next": "Save and continue", "skip_q": "Why are you skipping this test?", "choose": "Please choose …",
        "skip_reasons": ["I can't test this (e.g. I don't have what it needs)", "I don't know how to do this", "Other reason"],
        "next": "Next", "back": "← Back",
        "done_eyebrow": "Done", "done": "Thank you! All tests are done.",
        "done_sub": "You've really helped us. Now just save the result and send it.",
        "labels": {"ok": "works", "bad": "doesn't work", "skip": "skipped", "open": "open"},
        "device": "Finally, a few details about your device (optional):",
        "final": "Last step: save the result and send the file by e-mail to {email}.",
        "mail_btn": "Open your e-mail program", "mail_subject": "Test result {version} {platform}",
        "mail_body": "Hello,\n\nattached is my test result.\n\n(Please attach the file {file} from your Downloads folder.)",
        "mail_hint": "You attach the file yourself, it's in your Downloads folder. No e-mail program? Just send it from your webmail to {email}.",
        "save": "Save result", "saved": "Saved as \"{file}\" in your Downloads folder.",
        "back_last": "← Change the last test", "reset": "Start over",
        "reset_confirm": "Really delete all answers?", "reset_yes": "Yes, delete everything", "reset_no": "Cancel",
        "result_head": "Test result", "status": {"ok": "OK", "bad": "FAILED", "skip": "SKIPPED", "open": "OPEN", "na": "NOT APPLICABLE"},
    },
}


@dataclasses.dataclass
class Scenario:
    id: str
    title: str
    steps: list
    expected: str
    note: str
    only: set
    # step index -> (image name, caption); an image belongs to the step above it
    step_images: dict = dataclasses.field(default_factory=dict)
    result_image: tuple = None
    # Archived variant markers remain readable; build() applies STEP scenarios to everyone.
    variant: str = ""


@dataclasses.dataclass
class Chapter:
    title: str
    scenarios: list


# German and English scenario files use their own keywords.
FIELD_KEYWORDS = {
    "Nur": "only", "Only": "only",
    "Schritte": "steps", "Steps": "steps",
    "Erwartet": "expected", "Expected": "expected",
    "Hinweis": "note", "Note": "note",
    "Bild-Ergebnis": "result_image", "Result image": "result_image",
    "Variante": "variant", "Variant": "variant",
}
VARIANTS = ("standard", "step")
# "Bild:"/"Image:" lines under a step attach a screenshot to that step.
STEP_IMAGE_RE = re.compile(r"^\s+(?:Bild|Image):\s*(\S+)\s*(?:·\s*(.*?))?\s*$")
IMAGE_SPEC_RE = re.compile(r"^(\S+)\s*(?:·\s*(.*?))?\s*$")
CHAPTER_RE = re.compile(r"^#\s+(.+?)\s*$")
SCENARIO_RE = re.compile(r"^##\s+(\S+)\s*·\s*(.+?)\s*$")


def parse_scenarios(text):
    """Parse the scenario Markdown format documented in docs/tests/README.md."""
    chapters = []
    chapter = None
    scenario = None
    section = None  # which multi-line field ("Schritte") is currently being collected
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        m = CHAPTER_RE.match(line)
        if m:
            if scenario is not None:
                _finish_scenario(scenario)
            chapter = Chapter(title=m.group(1), scenarios=[])
            chapters.append(chapter)
            scenario = None
            section = None
            continue
        m = SCENARIO_RE.match(line)
        if m:
            if scenario is not None:
                _finish_scenario(scenario)
            if chapter is None:
                raise ValueError(f"Szenario {m.group(1)}: kein Kapitel (# …) davor")
            scenario = Scenario(id=m.group(1), title=m.group(2), steps=[], expected="", note="", only=set())
            chapter.scenarios.append(scenario)
            section = None
            continue
        if scenario is None:
            # Free text between a chapter heading and its first scenario would
            # never reach the guide, so it is an error rather than lost silently.
            if line.strip():
                raise ValueError(f"Text ohne Szenario wird nicht übernommen: {line.strip()[:60]}")
            continue
        img_m = STEP_IMAGE_RE.match(line)
        if img_m and section == "steps":
            if not scenario.steps:
                raise ValueError(f"Szenario {scenario.id}: Bild ohne Schritt davor")
            scenario.step_images[len(scenario.steps) - 1] = (img_m.group(1), img_m.group(2) or "")
            continue
        keyword, _, rest = line.partition(":")
        field = FIELD_KEYWORDS.get(keyword) if _ else None
        if field == "only":
            scenario.only = {p.strip() for p in rest.split(",") if p.strip()}
            section = None
        elif field == "steps":
            section = "steps"
        elif field == "expected":
            scenario.expected = rest.strip()
            section = None
        elif field == "note":
            scenario.note = rest.strip()
            section = None
        elif field == "variant":
            scenario.variant = rest.strip().lower()
            if scenario.variant not in VARIANTS:
                raise ValueError(f"Szenario {scenario.id}: Variante muss standard oder step sein")
            section = None
        elif field == "result_image":
            spec = IMAGE_SPEC_RE.match(rest.strip())
            if not spec:
                raise ValueError(f"Szenario {scenario.id}: Bild-Ergebnis ohne Namen")
            scenario.result_image = (spec.group(1), spec.group(2) or "")
            section = None
        elif section == "steps":
            step_m = re.match(r"^\d+\.\s*(.+)$", line)
            if step_m:
                scenario.steps.append(step_m.group(1).strip())
    if scenario is not None:
        _finish_scenario(scenario)
    return chapters


def _finish_scenario(scenario):
    if not scenario.expected or not scenario.steps:
        raise ValueError(f"Szenario {scenario.id}: Erwartet/Expected oder Schritte/Steps fehlen")


def attach_images(chapters, root, lang, platform):
    """Replace the image names of every scenario with data URIs. Screenshots
    live in docs/tests/bilder/ (German UI) and bilder/en/ (English UI); a
    platform-specific file `<name>.<platform>.webp` wins over `<name>.webp`,
    for pictures that show a path. A referenced picture that doesn't exist is
    an error, so a typo never ships a guide with a missing step picture."""
    folder = root / "docs/tests" / LANGS[lang]["bilder"]

    def uri(name, scenario_id):
        for candidate in (folder / f"{name}.{platform}.webp", folder / f"{name}.webp"):
            if candidate.is_file():
                return "data:image/webp;base64," + base64.b64encode(candidate.read_bytes()).decode("ascii")
        raise ValueError(f"Szenario {scenario_id}: Bild fehlt: {folder / name}.webp")

    for chapter in chapters:
        for s in chapter.scenarios:
            s.step_images = {i: (uri(name, s.id), cap) for i, (name, cap) in s.step_images.items()}
            if s.result_image:
                s.result_image = (uri(s.result_image[0], s.id), s.result_image[1])


def fill_scenario_text(chapters, values):
    """Fill {{name}} placeholders in scenario text (title, steps, expected,
    note) with the platform's values, in place. Lets a scenario that applies
    to all platforms name a path that actually differs per platform (e.g.
    {{datenordner}}) instead of forcing a platform-specific copy."""
    for chapter in chapters:
        for s in chapter.scenarios:
            s.title = fill(s.title, values)
            s.steps = [fill(step, values) for step in s.steps]
            s.expected = fill(s.expected, values)
            if s.note:
                s.note = fill(s.note, values)
            s.step_images = {i: (name, fill(cap, values)) for i, (name, cap) in s.step_images.items()}
            if s.result_image:
                s.result_image = (s.result_image[0], fill(s.result_image[1], values))


PLACEHOLDER_RE = re.compile(r"\{\{(\w+)\}\}")


def fill(template, values):
    """Substitute {{name}} placeholders; an unknown name is a hard error so a
    typo in a building block never silently ships an empty guide."""
    def repl(m):
        name = m.group(1)
        if name not in values:
            raise ValueError(f"Platzhalter unbekannt: {name}")
        return values[name]
    return PLACEHOLDER_RE.sub(repl, template)


def platform_values(version, platform, lang="de"):
    asset_platform, ext, path_template, log_template = PLATFORM_INFO[platform]
    name = "<your name>" if lang == "en" else "<dein Name>"
    datenordner = path_template.format(id=IDENTIFIER).replace("<dein Name>", name)
    logordner = log_template.format(id=IDENTIFIER).replace("<dein Name>", name)
    paket = release_assets.asset_name(version, asset_platform, ext, preview=True)
    return {
        "version": version,
        "produkt": PRODUCT,
        "datenordner": datenordner,
        "logordner": logordner,
        "paket": paket,
        "release": f"https://github.com/{REPO}/releases/tag/preview",
        "formular": f"{LANGS[lang]['form']}?version={version}&os={platform}",
        "testdaten": f"https://github.com/{REPO}/releases/download/preview/3MF-Testdaten.zip",
    }


UNRELEASED_RE = re.compile(r"## \[Unreleased\]\n(.*?)(?=\n## \[|\Z)", re.S)


def changelog_unreleased(text, lang="de"):
    """Return the '[Unreleased]' section in one language. CHANGELOG.md has the
    English text first and the German one after the '# Changelog (Deutsch)'
    heading; without that heading the first block is used for both."""
    marker = "# Changelog (Deutsch)"
    idx = text.find(marker)
    if idx == -1:
        search_text = text
    else:
        search_text = text[idx:] if lang == "de" else text[:idx]
    m = UNRELEASED_RE.search(search_text)
    return m.group(1).strip() if m else ""


def _markdown_inline(text):
    text = html.escape(text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)
    return text


def markdown_to_html(text):
    """Minimal Markdown->HTML for building blocks and the changelog excerpt:
    paragraphs, '-' lists, **bold**, `code`, [text](url) links."""
    lines = text.strip().splitlines()
    out = []
    para = []
    list_items = []

    def flush_para():
        if para:
            out.append(f"<p>{_markdown_inline(' '.join(para))}</p>")
            para.clear()

    def flush_list():
        if list_items:
            out.append("<ul>" + "".join(f"<li>{_markdown_inline(i)}</li>" for i in list_items) + "</ul>")
            list_items.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush_para()
            flush_list()
        elif stripped.startswith("#"):
            flush_para()
            flush_list()
            heading = stripped.lstrip("#").strip()
            out.append(f"<h3>{_markdown_inline(heading)}</h3>")
        elif stripped.startswith("- "):
            flush_para()
            list_items.append(stripped[2:].strip())
        else:
            flush_list()
            para.append(stripped)
    flush_para()
    flush_list()
    return "\n".join(out)


PAGE_CSS = """
@page { size: A4; margin: 18mm 16mm; }
body { font-family: "Segoe UI", system-ui, Arial, sans-serif; font-size: 10.5pt; line-height: 1.45; color: #222; }
h1 { font-size: 20pt; margin: 0 0 4pt; color: #c4502f; }
h2 { font-size: 14pt; margin: 20pt 0 6pt; color: #c4502f; border-bottom: 1.5px solid #f0c3b5; padding-bottom: 3pt; page-break-before: auto; page-break-after: avoid; }
h3 { font-size: 11.5pt; margin: 0 0 4pt; page-break-after: avoid; }
.sub { color: #666; margin-bottom: 12pt; }
code { font-family: Consolas, monospace; font-size: 9.5pt; background: #f4f1ee; padding: 1px 4px; border-radius: 3px; }
.szenario { page-break-inside: avoid; border: 1px solid #e5ddd6; border-radius: 6px; padding: 8px 11px; margin: 9pt 0; }
.szenario h3 { margin-top: 0; }
.lbl { font-weight: 600; color: #555; }
.ergebnis { color: #8a5a47; font-size: 9.5pt; margin-top: 6pt; }
.bogen { border-left: 3px solid #c4502f; background: #fbf1ec; padding: 6px 10px; margin: 8pt 0 12pt; }
ol, ul { margin: 3pt 0 3pt 18pt; padding: 0; }
li { margin: 2pt 0; }
.bild { margin: 4pt 0 6pt; page-break-inside: avoid; }
.bild img { max-width: 100%; max-height: 95mm; border: 1px solid #e5ddd6; border-radius: 4px; }
.bild figcaption { font-size: 9pt; font-weight: 600; color: #b3261e; }
.bild.ok figcaption { color: #2f7d4f; }
"""


def _scenario_html(s, lang="de"):
    # Steps/expected/note may contain a `command` (e.g. the PowerShell one-liner in
    # P1), so they go through the same inline Markdown renderer as building blocks
    # and the changelog, rather than a plain html.escape that would leave the
    # backticks in place and print literally instead of rendering as <code>.
    t = LANGS[lang]
    def figure(image, css=""):
        src, cap = image
        return (f'<figure class="bild{css}"><img src="{src}" alt="{html.escape(cap)}">'
                + (f"<figcaption>{_markdown_inline(cap)}</figcaption>" if cap else "") + "</figure>")

    steps_html = "<ol>" + "".join(
        f"<li>{_markdown_inline(step)}{figure(s.step_images[i]) if i in s.step_images else ''}</li>"
        for i, step in enumerate(s.steps)) + "</ol>"
    result_html = figure((s.result_image[0], f'{t["result_image"]}: {s.result_image[1]}' if s.result_image[1] else t["result_image"]), " ok") if s.result_image else ""
    note_html = f'<p><span class="lbl">{t["note"]}</span> {_markdown_inline(s.note)}</p>' if s.note else ""
    # No checkbox to tick here: a PDF can't be filled in, results go into the sheet.
    return (
        f'<div class="szenario"><h3>{html.escape(s.id)} · {_markdown_inline(s.title)}</h3>'
        f'<p><span class="lbl">{t["steps"]}</span></p>{steps_html}'
        f'<p><span class="lbl">{t["expected"]}</span> {_markdown_inline(s.expected)}</p>'
        f"{result_html}{note_html}"
        f'<div class="ergebnis">→ {html.escape(t["result_hint"].format(id=s.id))}</div>'
        "</div>"
    )


def build(version, platform, out_dir, root=".", pdf=False, lang="de", assistant_only=False):
    root = pathlib.Path(root)
    out_dir = pathlib.Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    t = LANGS[lang]
    platform_name = PLATFORM_NAMES[platform]
    values = platform_values(version, platform, lang)
    values["download"] = t["download_one"].format(**values)

    chapters = load_scenarios(root, version, lang)
    fill_scenario_text(chapters, values)

    bausteine_dir = root / "docs/tests" / t["bausteine"]

    def baustein(name, overrides=None):
        path = bausteine_dir / f"{name}.md"
        if not path.is_file():
            raise ValueError(f"Baustein fehlt: {path}")
        return markdown_to_html(fill(path.read_text(encoding="utf-8"), {**values, **(overrides or {})}))

    changelog_path = root / "CHANGELOG.md"
    changelog_html = (
        markdown_to_html(changelog_unreleased(changelog_path.read_text(encoding="utf-8"), lang))
        if changelog_path.is_file() else ""
    )

    chapters = platform_chapters(chapters, platform)
    # Keep parsing archived variant markers; only STEP-capable scenarios apply.
    for chapter in chapters:
        chapter.scenarios = [s for s in chapter.scenarios if s.variant != "standard"]
        for scenario in chapter.scenarios:
            scenario.variant = ""
    attach_images(chapters, root, lang, platform)
    chapters_html = [
        f"<h2>{html.escape(c.title)}</h2>" + "\n".join(_scenario_html(s, lang) for s in c.scenarios)
        for c in chapters
    ]

    sheet_name = t["sheet_file"].format(version=version, platform=platform_name)
    title = t["guide_title"].format(produkt=PRODUCT, version=version, platform=platform_name)
    new_h, install_h, data_h, report_h, uninstall_h = t["headings"]
    page = f"""<!doctype html><html lang="{lang}"><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>{PAGE_CSS}</style></head><body>
<h1>{html.escape(title)}</h1>
<p class="bogen">{html.escape(t["sheet_intro"].format(sheet=sheet_name))}</p>
<h2>{new_h}</h2>
{changelog_html}
<h2>{install_h}</h2>
{baustein(f"installieren-{platform}")}
<h2>{t["testdata_heading"]}</h2>
{baustein(f"testdaten-{platform}")}
<h2>{data_h}</h2>
{baustein("daten")}
{''.join(chapters_html)}
<h2>{report_h}</h2>
{baustein("fehler-melden")}
<h2>{uninstall_h}</h2>
{baustein(f"deinstallieren-{platform}")}
</body></html>
"""
    a = ASSISTANT_TEXT[lang]
    # The assistant puts a download button for the system where the
    # install text names the package; the page swaps in the real file name.
    assistant_path = build_assistant(out_dir, chapters, version, platform, lang, [
        (a["install"], baustein(f"installieren-{platform}", {"download": "__DOWNLOAD__", "paket": "__PAKET__"})),
        (t["testdata_heading"], baustein(f"testdaten-{platform}")),
    ], packages(values))
    # Testers only get the assistant; the guide and sheet remain for local use.
    if assistant_only:
        return assistant_path, None

    html_path = out_dir / (t["guide_file"].format(version=version, platform=platform_name) + ".html")
    html_path.write_text(page, encoding="utf-8")

    sheet_path = out_dir / sheet_name
    write_result_sheet(sheet_path, chapters, version, platform, lang)

    if pdf:
        render_pdf(html_path, html_path.with_suffix(".pdf"))
    return html_path, sheet_path


def packages(values):
    return {"name": values["paket"], "url": f"https://github.com/{REPO}/releases/download/preview/{values['paket']}"}


def build_assistant(out_dir, chapters, version, platform, lang, prep, packages):
    """Writes the interactive test assistant: a standalone HTML file that shows
    one test at a time, keeps its progress in the browser and saves the answers
    as a text file. `chapters` are already filtered for the platform and filled;
    `prep` lists the (title, html) screens shown before the first test."""
    a = ASSISTANT_TEXT[lang]
    platform_name = PLATFORM_NAMES[platform]
    # Scope of each test: the newest "new in" chapter is what returning testers
    # need; setup, basics, the bug report form and the update make the short test;
    # broken files, the lock test and log details go to an optional expert chapter.
    newest = next((m.group(1) for c in reversed(chapters) if (m := NEW_CHAPTER_RE.search(c.title))), None)

    def level(c, s):
        if s.id.startswith(FULL_ONLY_PREFIXES):
            return "full"
        found = NEW_CHAPTER_RE.search(c.title)
        if (found and found.group(1) == newest) or s.id.startswith("U"):
            return "new"
        if s.id in EXPERT_TESTS:
            return "expert"
        if s.id[0] in "EG" or s.id in ("F1", "F2"):
            return "short"
        return "full"

    tests = [
        {"id": s.id, "chapter": a["profi_chapter"] if level(c, s) == "expert" else c.title, "level": level(c, s),
         "title": _markdown_inline(s.title),
         "steps": [{"t": _markdown_inline(step),
                    "img": s.step_images[i][0] if i in s.step_images else "",
                    "cap": _markdown_inline(s.step_images[i][1]) if i in s.step_images else ""}
                   for i, step in enumerate(s.steps)],
         "expected": _markdown_inline(s.expected), "note": _markdown_inline(s.note) if s.note else "",
         "resultImg": s.result_image[0] if s.result_image else "",
         "resultCap": _markdown_inline(s.result_image[1]) if s.result_image else ""}
        for c in chapters for s in c.scenarios
    ]
    # Expert tests move to their own chapter just before the update tests.
    expert = [t for t in tests if t["level"] == "expert"]
    rest = [t for t in tests if t["level"] != "expert"]
    first_update = next((i for i, t in enumerate(rest) if t["id"].startswith("U")), len(rest))
    tests = rest[:first_update] + expert + rest[first_update:]
    # The assistant records the date automatically.
    date_field = LANGS[lang]["fields"]["common"][0]
    fields = [{"label": label, "hint": hint or "", "choices": list(choices or [])}
              for label, hint, choices in LANGS[lang]["fields"][platform] + LANGS[lang]["fields"]["common"]
              if label != date_field[0]]
    title = a["title"].format(produkt=PRODUCT, version=version, platform=platform_name)
    data = {
        "text": a, "tests": tests, "fields": fields, "package": packages, "prep": [{"title": ti, "html": h} for ti, h in prep], "title": title,
        "version": version, "platform": platform_name,
        "resultFile": a["result_file"].format(version=version, platform=platform_name),
        "email": TEST_EMAIL,
        "storageKey": f"3mf-testassistent-{version}-{platform}-{lang}",
    }
    # "</" inside the JSON would end the <script> element early.
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    page = (ASSISTANT_TEMPLATE.replace("__LANG__", lang).replace("__TITLE__", html.escape(title))
            .replace("__PRODUCT__", html.escape(PRODUCT)).replace("__VERSION__", html.escape(version))
            .replace("__PLATFORM__", html.escape(platform_name)).replace("__DATA__", payload)
            .replace("__FIGUREN__", FIGUREN_JS.read_text(encoding="utf-8").replace("</", "<\\/")))
    path = out_dir / a["file"].format(version=version, platform=platform_name)
    path.write_text(page, encoding="utf-8")
    return path


def load_scenarios(root, version, lang):
    """Scenarios of one language. The English file must list the same IDs in the
    same order as the German one, so both sheets and guides stay in step."""
    def read(suffix):
        path = root / "docs/tests" / f"{version}{suffix}"
        if not path.is_file():
            raise ValueError(f"Szenario-Datei fehlt: {path}")
        return parse_scenarios(path.read_text(encoding="utf-8"))

    chapters = read(LANGS[lang]["scenario_suffix"])
    if lang != "de":
        ids = [(s.id, sorted(s.only), s.variant) for c in chapters for s in c.scenarios]
        german = [(s.id, sorted(s.only), s.variant) for c in read(LANGS["de"]["scenario_suffix"]) for s in c.scenarios]
        if ids != german:
            raise ValueError(f"Szenarien ({lang}) passen nicht zur deutschen Datei: {ids} statt {german}")
    return chapters


def platform_chapters(chapters, platform):
    """Only the scenarios for this platform; a chapter left empty disappears."""
    result = []
    for c in chapters:
        scenarios = [s for s in c.scenarios if not s.only or platform in s.only]
        if scenarios:
            result.append(Chapter(title=c.title, scenarios=scenarios))
    return result


SHEET_COLORS = {"accent": "FF7A5C", "ink": "1F1B18", "muted": "6D665F", "hint": "A39B92",
                "chapter": "EFEBE7", "input": "FFF8E6", "line": "D9D4CE"}
RESULT_FILLS = ("D8F3E3", "FDDCDC", "ECEAE7")  # OK, failed, skipped


def write_result_sheet(path, chapters, version, platform, lang):
    """Result sheet for one platform: device fields, one row per scenario with a
    result list (OK/failed/skipped, coloured) and a note, and a summary."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    t = LANGS[lang]
    c = SHEET_COLORS
    fill = lambda color: PatternFill("solid", start_color=color, end_color=color)
    side = Side(style="thin", color=c["line"])
    box = Border(left=side, right=side, top=side, bottom=side)
    results = t["results"]

    wb = Workbook()
    ws = wb.active
    ws.title = t["sheet_tab"]
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCD", (7, 44, 18, 52)):
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:D1")
    ws["A1"] = t["sheet_title"].format(produkt=PRODUCT, version=version, platform=PLATFORM_NAMES[platform])
    ws["A1"].font = Font(size=16, bold=True, color=c["ink"])
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:D2")
    ws["A2"] = t["sheet_help"]
    ws["A2"].font = Font(size=10, color=c["muted"])
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30

    row = 4
    for label, hint, choices in t["fields"][platform] + t["fields"]["common"]:
        ws.cell(row, 1, label).font = Font(bold=True, color=c["ink"])
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        for col in (3, 4):
            ws.cell(row, col).fill = fill(c["input"])
            ws.cell(row, col).border = box
        ws.row_dimensions[row].height = 20
        value = ws.cell(row, 3)
        value.alignment = Alignment(vertical="center")
        if hint:
            value.value = hint
            value.font = Font(italic=True, color=c["hint"])
        if choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(value)
        row += 1

    row += 1
    header_row = row
    for col, text in enumerate(t["columns"], start=1):
        cell = ws.cell(row, col, text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill(c["accent"])
        cell.alignment = Alignment(vertical="center")
        cell.border = box
    ws.row_dimensions[row].height = 22
    row += 1

    result_list = DataValidation(type="list", formula1='"' + ",".join(results) + '"', allow_blank=True,
                                 showErrorMessage=True, errorTitle=t["columns"][2], error=t["invalid_result"])
    ws.add_data_validation(result_list)
    first = row
    for chapter in chapters:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row, 1, chapter.title)
        cell.font = Font(bold=True, color=c["ink"])
        cell.fill = fill(c["chapter"])
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 20
        row += 1
        for s in chapter.scenarios:
            ws.cell(row, 1, s.id).font = Font(bold=True, color=c["muted"])
            # Titles may carry inline Markdown (`code`); the sheet shows plain text.
            ws.cell(row, 2, s.title.replace("`", ""))
            result_list.add(ws.cell(row, 3))
            for col in range(1, 5):
                ws.cell(row, col).border = box
                ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[row].height = 30
            row += 1
    last = row - 1

    results_range = f"C{first}:C{last}"
    for value, color in zip(results, RESULT_FILLS):
        ws.conditional_formatting.add(
            results_range, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill(color)))

    row += 1
    ws.cell(row, 2, t["summary"]).font = Font(bold=True, color=c["ink"])
    row += 1
    # Chapter rows only fill column A, so counting column B counts scenarios.
    counts = [(value, f'=COUNTIF({results_range},"{value}")') for value in results]
    counts.append((t["open"], f"=COUNTA(B{first}:B{last})-COUNTA({results_range})"))
    for label, formula in counts:
        ws.cell(row, 2, label)
        ws.cell(row, 3, formula).alignment = Alignment(horizontal="left")
        row += 1

    # Calibri (openpyxl's default) is missing in Numbers and LibreOffice on Linux.
    for sheet_row in ws.iter_rows():
        for cell in sheet_row:
            font = copy.copy(cell.font)
            font.name = "Arial"
            cell.font = font

    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)


# The website's mascots (copy of 3mf-katalog-webseite/figuren.js, keep in sync).
FIGUREN_JS = pathlib.Path(__file__).resolve().parent / "figuren.js"

# Standalone page of the test assistant. Offline by design: no external fonts
# or scripts, so it works from a double-click in the Downloads folder.
ASSISTANT_TEMPLATE = """<!doctype html>
<html lang="__LANG__"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITLE__</title>
<style>
:root{--bg:#f6f2ef;--surface:#fff;--ink:#2b2320;--muted:#7a6a62;--line:#e5ddd6;--accent:#c4502f;--accent-ink:#fff;--accent-soft:#fbf1ec;
--ok:#2f7d4f;--ok-soft:#e6f3eb;--bad:#b3261e;--bad-soft:#fbe9e7;--skip:#8a7a3c;--skip-soft:#f6f1dc;--code:#f4f1ee;color-scheme:light}
@media (prefers-color-scheme:dark){:root{--bg:#1d1917;--surface:#282220;--ink:#f1e9e4;--muted:#b3a39a;--line:#3d3431;--accent:#e0714f;--accent-ink:#1d1917;
--accent-soft:#3a2721;--ok:#6cc08b;--ok-soft:#1f3326;--bad:#f08a80;--bad-soft:#3d2220;--skip:#d6c27a;--skip-soft:#353019;--code:#332b28;color-scheme:dark}}
*{box-sizing:border-box}
[hidden]{display:none!important}
body{margin:0;background:var(--bg);color:var(--ink);font-family:"Segoe UI",system-ui,-apple-system,Arial,sans-serif;font-size:17px;line-height:1.5;padding:20px 16px 48px}
.wrap{max-width:720px;margin:0 auto;display:flex;flex-direction:column;gap:16px}
header{display:flex;align-items:center;justify-content:space-between;gap:12px;flex-wrap:wrap}
.brand{font-weight:700}.brand span{color:var(--muted);font-weight:400}
.platform{font-size:13px;padding:3px 10px;border-radius:99px;background:var(--accent-soft);color:var(--accent);font-weight:600;letter-spacing:.03em;text-transform:uppercase}
.progress{display:flex;flex-direction:column;gap:6px}.progress .row{display:flex;justify-content:space-between;font-size:14px;color:var(--muted);font-variant-numeric:tabular-nums}
.shell{display:flex;justify-content:center;align-items:flex-start;gap:24px}
.shell>.wrap{flex:1 1 720px;max-width:720px;min-width:0;margin:0}
header .left{display:flex;align-items:center;gap:10px;min-width:0}
.burger{padding:6px 10px;font-size:18px;line-height:1;border-radius:8px}
.chaps{background:var(--surface);border:1px solid var(--line);border-radius:14px;padding:14px 10px;display:flex;flex-direction:column;gap:4px}
.chaps.side{position:sticky;top:20px;flex:0 0 260px;max-height:calc(100vh - 40px);overflow:auto}
.chaps.drawer{position:fixed;top:0;left:0;bottom:0;width:min(320px,86vw);z-index:20;border-radius:0 14px 14px 0;overflow:auto;padding-top:18px;box-shadow:0 10px 40px rgba(0,0,0,.35)}
.backdrop{position:fixed;inset:0;background:rgba(0,0,0,.45);z-index:19}
.chaps .lbl{margin:0 6px 2px}.chaps .tip{margin:0 6px 8px;font-size:12.5px}
.chaps button{display:flex;flex-direction:column;align-items:stretch;gap:5px;text-align:left;padding:8px 10px;font-weight:600;font-size:14.5px;border-width:1px;border-color:transparent;background:none;border-radius:8px}
.chaps button:hover{background:var(--bg)}
.chaps button .top{display:flex;justify-content:space-between;gap:8px;align-items:baseline}
.chaps button .top span{font-weight:400;font-size:12.5px;color:var(--muted);white-space:nowrap;font-variant-numeric:tabular-nums}
.chaps button .mini{height:4px;border-radius:99px;background:var(--line);overflow:hidden}.chaps button .mini i{display:block;height:100%;background:var(--accent)}
.chaps button.cur{border-color:var(--accent);background:var(--accent-soft)}
.chaps button.full .top span{color:var(--ok)}.chaps button.full .mini i{background:var(--ok)}
.bar{height:8px;border-radius:99px;background:var(--line);overflow:hidden}.bar i{display:block;height:100%;background:var(--accent);border-radius:99px;transition:width .3s}
.card{background:var(--surface);border:1px solid var(--line);border-radius:14px;padding:24px;display:flex;flex-direction:column;gap:18px}
.eyebrow{font-size:13px;text-transform:uppercase;letter-spacing:.06em;color:var(--muted);font-weight:600}
h1{font-size:26px;line-height:1.25;margin:0}h1 .id{color:var(--accent);margin-right:6px}
.card p{margin:0}
.hero{display:flex;align-items:center;gap:22px}
.hero .fig{width:118px;height:134px;flex:none;overflow:visible;filter:drop-shadow(0 10px 16px rgba(60,40,30,.28))}
.hero h1{margin-bottom:6px}.hero p{color:var(--muted)}
.chip{display:inline-block;white-space:nowrap;font-size:14px;font-weight:700;padding:3px 10px;border-radius:99px;background:var(--accent-soft);color:var(--accent);margin-top:10px}
@media (max-width:520px){.hero{flex-direction:column;text-align:center}}
@media (prefers-color-scheme:dark){.hero .fig{filter:none}}
.tiles{display:grid;grid-template-columns:1fr 1fr;gap:12px}@media (max-width:560px){.tiles{grid-template-columns:1fr}}
.tile{border:1px solid var(--line);border-radius:12px;padding:14px 16px;background:var(--bg);display:flex;flex-direction:column;gap:4px}
.tile b{font-size:16px}.tile span{font-size:15px;color:var(--muted);line-height:1.45}
.section{border-top:1px solid var(--line);padding-top:18px;display:flex;flex-direction:column;gap:12px}
.pose-point .arm{transform-box:fill-box;transform-origin:center;animation:poke 2.4s ease-in-out infinite}
@keyframes poke{0%,70%,100%{transform:scale(1)}78%{transform:scale(1.14)}86%{transform:scale(1)}92%{transform:scale(1.1)}}
.pose-thumbs .arm{transform-box:fill-box;transform-origin:center;animation:wipp 2.6s ease-in-out infinite}
@keyframes wipp{0%,70%,100%{transform:translateY(0)}78%{transform:translateY(-6px)}86%{transform:translateY(0)}92%{transform:translateY(-4px)}}
.fig .blink{animation:blink 4s infinite;transform-box:fill-box;transform-origin:center}
@keyframes blink{0%,94%,100%{transform:scaleY(1)}96%{transform:scaleY(.1)}}
@media (prefers-reduced-motion:reduce){.fig .arm,.fig .blink{animation:none}}.card ul{margin:0;padding-left:20px;display:flex;flex-direction:column;gap:6px}
ol.steps{margin:0;padding:0;list-style:none;display:flex;flex-direction:column;gap:10px;counter-reset:s}
ol.steps li{counter-increment:s;display:grid;grid-template-columns:34px 1fr;gap:12px;align-items:start;cursor:pointer;padding:8px;border-radius:10px}
ol.steps li:hover{background:var(--bg)}
ol.steps li::before{content:counter(s);width:30px;height:30px;border-radius:50%;display:grid;place-items:center;font-weight:700;font-size:15px;background:var(--accent-soft);color:var(--accent)}
ol.steps li.done{color:var(--muted)}ol.steps li.done::before{content:"\\2713";background:var(--ok-soft);color:var(--ok)}
.tip{font-size:13px;color:var(--muted);margin-top:-8px}
.stext{display:flex;flex-direction:column;gap:10px;min-width:0}
.shot{margin:0;display:flex;flex-direction:column;gap:6px}
.shot img{display:block;max-width:100%;height:auto;border:1px solid var(--line);border-radius:8px;cursor:zoom-in;background:#fff}
.shot figcaption{font-size:14px;font-weight:600;color:var(--bad)}
.shot.ok{margin-top:10px}.shot.ok figcaption{color:var(--ok)}
ol.steps li.done img{opacity:.55}
.lightbox{position:fixed;inset:0;background:rgba(0,0,0,.82);display:flex;flex-direction:column;align-items:center;justify-content:center;gap:10px;padding:16px;z-index:10;cursor:zoom-out}
.lightbox img{max-width:100%;max-height:88vh;border-radius:8px;background:#fff}.lightbox span{color:#fff;font-size:14px}
.expect{background:var(--accent-soft);border-left:4px solid var(--accent);border-radius:8px;padding:14px 16px}
.lbl{font-size:13px;font-weight:700;text-transform:uppercase;letter-spacing:.05em;color:var(--accent);display:block;margin-bottom:4px}
.note{border:1px solid var(--line);border-radius:8px;padding:12px 16px;font-size:15px}.note .lbl{color:var(--muted)}
.install h3{font-size:18px;margin:0}
.download{display:flex;flex-direction:column;align-items:center;gap:4px;padding:16px;border-radius:12px;background:var(--accent);color:var(--accent-ink);font-weight:700;font-size:19px;text-decoration:none;text-align:center}
.download span{font-weight:400;font-size:13px;opacity:.9;overflow-wrap:anywhere}
.download:hover{filter:brightness(1.08)}
.mailbtn{display:flex;flex-direction:column;align-items:center;gap:2px;padding:12px;border-radius:10px;border:2px solid var(--accent);color:var(--accent);font-weight:700;text-decoration:none;text-align:center}
.mailbtn span{font-weight:400;font-size:14px}.mailbtn:hover{background:var(--accent-soft)}
.dl{display:flex;flex-direction:column;gap:6px}
code{font-family:Consolas,ui-monospace,monospace;background:var(--code);padding:1px 5px;border-radius:4px;font-size:.92em;overflow-wrap:anywhere}
a{color:var(--accent)}
.q{font-weight:600}
.actions{display:grid;grid-template-columns:repeat(3,1fr);gap:10px}@media (max-width:520px){.actions{grid-template-columns:1fr}}
button{font:inherit;cursor:pointer;border-radius:10px;border:2px solid var(--line);background:var(--surface);color:var(--ink);padding:14px 12px;font-weight:600}
button:focus-visible,textarea:focus-visible,select:focus-visible,input:focus-visible,li:focus-visible{outline:3px solid var(--accent);outline-offset:2px}
.b-ok{border-color:var(--ok);color:var(--ok)}.b-ok:hover,.b-ok.sel{background:var(--ok-soft)}
.b-bad{border-color:var(--bad);color:var(--bad)}.b-bad:hover,.b-bad.sel{background:var(--bad-soft)}
.b-skip{border-color:var(--skip);color:var(--skip)}.b-skip:hover,.b-skip.sel{background:var(--skip-soft)}
.primary{background:var(--accent);border-color:var(--accent);color:var(--accent-ink)}.primary:disabled{opacity:.45;cursor:not-allowed}
.followup{display:flex;flex-direction:column;gap:10px;border-top:1px solid var(--line);padding-top:16px}
textarea,select,input{font:inherit;color:var(--ink);background:var(--bg);border:1px solid var(--line);border-radius:8px;padding:10px 12px;width:100%}
textarea{min-height:96px;resize:vertical}
.nav{display:flex;justify-content:space-between;align-items:center;gap:10px;flex-wrap:wrap}
.confirm{display:flex;flex-direction:column;align-items:flex-end;gap:8px}
.confirm-btns{display:flex;gap:10px;flex-wrap:wrap;justify-content:flex-end}
.confirm button{padding:8px 14px;font-size:15px}
.b-yes{border-color:var(--bad);background:var(--bad);color:var(--surface)}.b-yes:hover{filter:brightness(1.1)}
.b-no{border-color:var(--ok);background:var(--ok);color:var(--surface)}.b-no:hover{filter:brightness(1.1)}
.linkbtn{background:none;border:none;padding:6px 0;color:var(--muted);font-weight:500;text-decoration:underline;text-underline-offset:3px}
.sum{display:grid;grid-template-columns:repeat(3,1fr);gap:10px}.sum div{border-radius:10px;padding:12px;text-align:center}.sum b{display:block;font-size:28px;line-height:1.1}
.s-ok{background:var(--ok-soft);color:var(--ok)}.s-bad{background:var(--bad-soft);color:var(--bad)}.s-skip{background:var(--skip-soft);color:var(--skip)}
table{width:100%;border-collapse:collapse;font-size:15px}td{padding:8px 6px;border-top:1px solid var(--line);vertical-align:top}td:first-child{white-space:nowrap;font-weight:600}
.tag{font-size:12px;font-weight:700;padding:2px 8px;border-radius:99px;white-space:nowrap}
.t-ok{background:var(--ok-soft);color:var(--ok)}.t-bad{background:var(--bad-soft);color:var(--bad)}.t-skip{background:var(--skip-soft);color:var(--skip)}.t-open{background:var(--bg);color:var(--muted)}
.modes{display:grid;grid-template-columns:repeat(3,1fr);gap:10px}@media (max-width:620px){.modes{grid-template-columns:1fr}}
.modes button{display:flex;flex-direction:column;align-items:flex-start;gap:4px;text-align:left}
.modes button span{font-weight:400;font-size:14px;color:var(--muted);line-height:1.35}
.modes button em{font-style:normal;font-size:13px;color:var(--accent);font-variant-numeric:tabular-nums}
.modes button.sel{border-color:var(--accent);background:var(--accent-soft)}
.fields{display:grid;gap:10px}.fields label{font-size:14px;color:var(--muted);display:flex;flex-direction:column;gap:4px}
.saved{color:var(--ok);font-weight:600}
@media (prefers-reduced-motion:reduce){.bar i{transition:none}}
</style></head><body>
<div class="shell">
<nav class="chaps" id="chaps" hidden></nav>
<div class="wrap">
<header><div class="left"><button class="burger" id="chapBtn" hidden aria-expanded="false" aria-controls="chaps"></button><div class="brand">__PRODUCT__ <span>· __VERSION__</span></div></div><div class="platform">__PLATFORM__</div></header>
<div class="progress" id="progress" hidden><div class="row"><span id="progText"></span><span id="progCounts"></span></div><div class="bar"><i id="barFill"></i></div></div>
<main id="view"></main>
</div>
</div>
<div class="backdrop" id="chapBack" hidden></div>
<script id="data" type="application/json">__DATA__</script>
<script>__FIGUREN__</script>
<script>
const D = JSON.parse(document.getElementById("data").textContent), T = D.text;
["final", "mail_hint"].forEach(k => T[k] = T[k].split("{email}").join(D.email));
T.intro = T.intro.map(x => [x[0], x[1].split("{email}").join(D.email)]);
// One of the website's mascots (figuren.js), picked at random like on 3mfkatalog.de.
const FIGURE = window.MMKFiguren ? window.MMKFiguren.random() : null;
const mascot = pose => FIGURE ? '<span class="pose-' + pose + '">' + window.MMKFiguren.svg(FIGURE, pose) + '</span>' : '';
const fresh = () => ({ started:false, prepDone:false, prepStep:0, index:0, results:{}, done:{}, device:{}, mode:"", secs:{}, shown:null });
let state = Object.assign(fresh(), load() || {});
if (!state.mode && !state.started && returningTester()) state.mode = "new";
let TESTS = [];
// The selected test scope determines which scenarios are shown.
const IN_MODE = { new:["new"], short:["new","short"], full:["new","short","full","expert"] };
const inMode = (t, mode) => (IN_MODE[mode] || IN_MODE.full).includes(t.level);
function applyMode(){ TESTS = D.tests.filter(t => inMode(t, state.mode)); }
// A tester who already has progress stored for another test version is a returning tester.
function returningTester(){
  try { for (let i = 0; i < localStorage.length; i++) { const k = localStorage.key(i); if (k && k.startsWith("3mf-testassistent-") && k !== D.storageKey) return true; } } catch (e) {}
  return false;
}
// Update tests (U…) only happen once the next test version is out, so they are
// left out of the count and time shown when choosing the scope.
const modeCount = m => D.tests.filter(t => !t.id.startsWith("U") && inMode(t, m)).length;
applyMode();
const MAX_MS = 30 * 60 * 1000;
// Time on the current test, capped so a break or an open laptop overnight doesn't count.
function track(t){
  const now = Date.now(), s = state.shown && state.shown.id === t.id ? Math.min(now - state.shown.at, MAX_MS) : 0;
  state.secs[t.id] = (state.secs[t.id] || 0) + s; state.shown = { id:t.id, at:now };
}
const mins = ms => Math.round(ms / 60000);
const oneTime = ms => ms < 60000 ? T.time_lt1 : T.time_one.replace("{min}", mins(ms));
const totalMs = () => TESTS.reduce((a, t) => a + (state.secs[t.id] || 0), 0);
function load(){ try { return JSON.parse(localStorage.getItem(D.storageKey)); } catch(e) { return null; } }
function save(){ try { localStorage.setItem(D.storageKey, JSON.stringify(state)); } catch(e) {} }
const view = document.getElementById("view");
const esc = s => String(s).replace(/[&<>"]/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
const counts = () => { const r = Object.values(state.results); return { ok:r.filter(x=>x.s==="ok").length, bad:r.filter(x=>x.s==="bad").length, skip:r.filter(x=>x.s==="skip").length }; };
function progress(){
  const p = document.getElementById("progress"); p.hidden = !state.started || !state.prepDone; if (p.hidden) return;
  const i = Math.min(state.index, TESTS.length), c = counts();
  document.getElementById("progText").textContent = i >= TESTS.length ? T.all_done : T.test_of.replace("{n}", i+1).replace("{total}", TESTS.length);
  document.getElementById("progCounts").textContent = "✅ " + c.ok + " · ❌ " + c.bad + " · ⏭ " + c.skip;
  document.getElementById("barFill").style.width = (Object.keys(state.results).length / TESTS.length * 100) + "%";
}
// Chapters in test order, so a tester can jump straight to e.g. the new tests
// of this version. Answers stay where they are; jumped-over tests stay open.
function chapterList(){
  const out = [];
  TESTS.forEach((t, i) => { let c = out[out.length - 1]; if (!c || c.title !== t.chapter) out.push(c = { title:t.chapter, first:i, ids:[] }); c.ids.push(t.id); });
  return out;
}
// Fixed list left of the test on wide screens; on narrow screens a ☰ button
// opens the same list as a drawer from the left.
const wide = window.matchMedia("(min-width: 1060px)");
let navOpen = false;
function renderChapters(){
  const nav = document.getElementById("chaps"), btn = document.getElementById("chapBtn"), back = document.getElementById("chapBack");
  const avail = state.started && state.prepDone, side = wide.matches;
  if (!avail || side) navOpen = false;
  btn.hidden = !avail || side; btn.textContent = "☰"; btn.title = T.chapters_title; btn.setAttribute("aria-label", T.chapters_title); btn.setAttribute("aria-expanded", String(navOpen));
  nav.hidden = !avail || !(side || navOpen); back.hidden = !navOpen;
  nav.className = "chaps " + (side ? "side" : "drawer");
  if (nav.hidden) return;
  const cur = TESTS[state.index] ? TESTS[state.index].chapter : null, list = chapterList();
  nav.innerHTML = '<span class="lbl">' + esc(T.chapters_btn) + '</span><p class="tip">' + esc(T.chapters_hint) + '</p>'
    + list.map((c, k) => { const done = c.ids.filter(id => state.results[id]).length;
        return '<button data-k="' + k + '" class="' + (c.title === cur ? "cur" : "") + (done === c.ids.length ? " full" : "") + '"' + (c.title === cur ? ' aria-current="true"' : '') + '><span class="top">' + esc(c.title)
          + '<span>' + done + '/' + c.ids.length + '</span></span><span class="mini"><i style="width:' + (done / c.ids.length * 100) + '%"></i></span></button>'; }).join("");
  nav.querySelectorAll("button").forEach(b => b.onclick = () => {
    const c = list[+b.dataset.k], open = c.ids.findIndex(id => !state.results[id]);
    state.index = c.first + (open < 0 ? 0 : open); navOpen = false; save(); render(); window.scrollTo(0, 0);
  });
}
function setNav(open){ navOpen = open; renderChapters(); if (open) document.querySelector("#chaps button")?.focus(); else document.getElementById("chapBtn").focus(); }
document.getElementById("chapBtn").onclick = () => setNav(!navOpen);
document.getElementById("chapBack").onclick = () => setNav(false);
document.addEventListener("keydown", e => { if (e.key === "Escape" && navOpen) setNav(false); });
wide.addEventListener("change", renderChapters);
function render(){
  progress();
  renderChapters();
  if (!state.started) return renderStart();
  if (!state.prepDone) return renderPrep();
  if (state.index >= TESTS.length) return renderSummary();
  renderTest(TESTS[state.index]);
}
function renderStart(){
  const resume = Object.keys(state.results).length > 0 || state.prepDone;
  view.innerHTML = '<section class="card"><div class="hero">' + mascot("point") + '<div><div class="eyebrow">' + esc(T.eyebrow) + '</div><h1>' + esc(T.hello) + '</h1>'
    + '<p>' + esc(T.sub) + '</p><span class="chip">' + esc(D.version) + ' · ' + esc(D.platform) + '</span></div></div>'
    + '<div class="tiles">' + T.intro.map(x => '<div class="tile"><b>' + esc(x[0]) + '</b><span>' + esc(x[1]) + '</span></div>').join("") + '</div>'
    + '<div class="section"><p class="q">' + esc(T.mode_q) + '</p>'
    + (returningTester() ? '<div class="tip" style="margin-top:0">' + esc(T.mode_returning) + '</div>' : '')
    + '<div class="modes">' + ["new", "short", "full"].map(m => '<button data-m="' + m + '" class="' + (state.mode === m ? "sel" : "") + '"><b>' + esc(T.modes[m][0]) + '</b><span>' + esc(T.modes[m][1]) + '</span><em>'
        + esc(T.mode_meta.replace("{n}", modeCount(m)).replace("{min}", Math.max(5, Math.round(modeCount(m) * 3.5 / 5) * 5))) + '</em></button>').join("") + '</div><div class="tip" style="margin-top:0">' + esc(T.mode_update_note) + '</div></div>'
    + '<div class="section"><button class="primary" id="go"' + (state.mode ? '' : ' disabled') + '>' + esc(resume ? T.resume : T.start) + '</button></div></section>';
  view.querySelectorAll(".modes button").forEach(b => b.onclick = () => { state.mode = b.dataset.m; applyMode(); if (state.index > TESTS.length) state.index = TESTS.length; save(); renderStart(); });
  document.getElementById("go").onclick = () => { state.started = true; save(); render(); };
}
function renderPrep(){
  const i = Math.min(state.prepStep || 0, D.prep.length - 1), p = D.prep[i];
  const pkg = D.package;
  const button = '<a class="download" href="' + esc(pkg.url) + '" download>⬇ ' + esc(T.download_btn) + '<span>' + esc(pkg.name) + '</span></a><div class="tip" style="margin-top:0">' + esc(T.download_note) + '</div>';
  const body = p.html.replace(/<p>__DOWNLOAD__/, '<div class="dl">' + button + '</div><p>').replace(/__PAKET__/g, esc(pkg.name));
  view.innerHTML = '<section class="card install"><div class="eyebrow">' + esc(T.prep) + ' ' + (i + 1) + '/' + D.prep.length + '</div><h1>' + esc(p.title) + '</h1>' + body
    + '<button class="primary" id="prepOk">' + esc(T.prep_done) + '</button>'
    + '<div class="nav">' + (i > 0 ? '<button class="linkbtn" id="prepBack">' + esc(T.back) + '</button>' : '<span></span>') + resetBox() + '</div></section>';
  view.querySelectorAll("a").forEach(a => { a.target = "_blank"; a.rel = "noopener"; });
  document.getElementById("prepOk").onclick = () => { state.prepStep = i + 1; if (state.prepStep >= D.prep.length) state.prepDone = true; save(); render(); window.scrollTo(0, 0); };
  if (i > 0) document.getElementById("prepBack").onclick = () => { state.prepStep = i - 1; save(); render(); };
  wireReset();
}
function renderTest(t){
  if (!state.shown || state.shown.id !== t.id) { state.shown = { id:t.id, at:Date.now() }; save(); }
  const r = state.results[t.id], doneSteps = state.done[t.id] || [];
  view.innerHTML = '<section class="card"><div class="eyebrow">' + esc(t.chapter) + '</div><h1><span class="id">' + esc(t.id) + '</span>' + t.title + '</h1>'
    + '<ol class="steps">' + t.steps.map((s, i) => '<li tabindex="0" data-i="' + i + '" class="' + (doneSteps.includes(i) ? "done" : "") + '"><div class="stext">' + s.t
        + (s.img ? '<figure class="shot"><img src="' + s.img + '" alt=""><figcaption>' + s.cap + '</figcaption></figure>' : '') + '</div></li>').join("") + '</ol>'
    + '<div class="tip">' + esc(T.tip) + '</div>'
    + '<div class="expect"><span class="lbl">' + esc(T.expected) + '</span>' + t.expected
        + (t.resultImg ? '<figure class="shot ok"><img src="' + t.resultImg + '" alt=""><figcaption>\u2705 ' + esc(T.result_image) + (t.resultCap ? ': ' + t.resultCap : '') + '</figcaption></figure>' : '') + '</div>'
    + (t.note ? '<div class="note"><span class="lbl">' + esc(T.note) + '</span>' + t.note + '</div>' : '')
    + '<p class="q">' + esc(T.question) + '</p><div class="actions">'
    + '<button class="b-ok" data-s="ok">✅ ' + esc(T.ok) + '</button><button class="b-bad" data-s="bad">❌ ' + esc(T.bad) + '</button><button class="b-skip" data-s="skip">⏭ ' + esc(T.skip) + '</button></div>'
    + '<div id="follow"></div><div class="nav"><button class="linkbtn" id="back">' + esc(T.back) + '</button>' + resetBox() + '</div></section>';
  view.querySelectorAll("a").forEach(a => { a.target = "_blank"; a.rel = "noopener"; });
  view.querySelectorAll("ol.steps li").forEach(li => {
    const toggle = () => { const i = +li.dataset.i, d = state.done[t.id] || (state.done[t.id] = []), k = d.indexOf(i); k < 0 ? d.push(i) : d.splice(k, 1); li.classList.toggle("done"); save(); };
    li.onclick = e => { if (e.target.closest("a")) return; if (e.target.tagName === "IMG") { zoom(e.target.src); return; } toggle(); };
    li.onkeydown = e => { if (e.key === "Enter" || e.key === " ") { e.preventDefault(); toggle(); } };
  });
  view.querySelectorAll(".expect img").forEach(im => im.onclick = () => zoom(im.src));
  // Screenshots are taken at 1.5x pixel density: show them at their real size
  // instead of stretching small crops to the full width.
  view.querySelectorAll(".shot img").forEach(im => {
    const fit = () => { if (im.naturalWidth) im.style.width = Math.round(im.naturalWidth / 1.5) + "px"; };
    im.complete ? fit() : im.addEventListener("load", fit);
  });
  view.querySelectorAll(".actions button").forEach(b => b.onclick = () => choose(t, b.dataset.s));
  document.getElementById("back").onclick = () => {
    if (state.index === 0) { state.prepDone = false; state.prepStep = D.prep.length - 1; } else state.index--;
    save(); render(); window.scrollTo(0, 0);
  };
  wireReset();
  if (r) { view.querySelector('.actions [data-s="' + r.s + '"]').classList.add("sel"); if (r.s !== "ok") choose(t, r.s, true); }
}
function choose(t, s, restoring){
  view.querySelectorAll(".actions button").forEach(b => b.classList.toggle("sel", b.dataset.s === s));
  const f = document.getElementById("follow"), prev = state.results[t.id] && state.results[t.id].s === s ? state.results[t.id] : null;
  if (s === "ok") { track(t); state.results[t.id] = { s:"ok" }; next(); return; }
  if (s === "bad") {
    f.innerHTML = '<div class="followup"><label class="q" for="why">' + esc(T.why) + '</label><textarea id="why" placeholder="' + esc(T.why_ph) + '">' + esc(prev ? prev.text : "") + '</textarea><button class="primary" id="cont" disabled>' + esc(T.save_next) + '</button></div>';
    const ta = document.getElementById("why"), c = document.getElementById("cont"), upd = () => c.disabled = ta.value.trim().length < 5;
    ta.oninput = upd; upd(); c.onclick = () => { track(t); state.results[t.id] = { s:"bad", text:ta.value.trim() }; next(); };
    if (!restoring) ta.focus();
  } else {
    f.innerHTML = '<div class="followup"><label class="q" for="reason">' + esc(T.skip_q) + '</label><select id="reason"><option value="">' + esc(T.choose) + '</option>'
      + T.skip_reasons.map(x => '<option' + (prev && prev.text === x ? ' selected' : '') + '>' + esc(x) + '</option>').join("") + '</select><button class="primary" id="cont" disabled>' + esc(T.next) + '</button></div>';
    const sel = document.getElementById("reason"), c = document.getElementById("cont"), upd = () => c.disabled = !sel.value;
    sel.onchange = upd; upd(); c.onclick = () => { track(t); state.results[t.id] = { s:"skip", text:sel.value }; next(); };
    if (!restoring) sel.focus();
  }
}
// "Start over" on every screen after the start page, always behind a confirmation.
const resetBox = () => '<span id="resetBox"><button class="linkbtn" id="reset">' + esc(T.reset) + '</button></span>';
function wireReset(){
  const r = document.getElementById("reset"); if (!r) return;
  r.onclick = () => {
    const box = document.getElementById("resetBox");
    box.innerHTML = '<span class="confirm"><span class="q">' + esc(T.reset_confirm) + '</span><span class="confirm-btns">'
      + '<button class="b-yes" id="resetYes">' + esc(T.reset_yes) + '</button><button class="b-no" id="resetNo">' + esc(T.reset_no) + '</button></span></span>';
    document.getElementById("resetNo").focus();
    document.getElementById("resetYes").onclick = () => { try { localStorage.removeItem(D.storageKey); } catch (e) {} state = fresh(); applyMode(); save(); render(); window.scrollTo(0, 0); };
    document.getElementById("resetNo").onclick = () => { box.innerHTML = resetBox().replace(/^<span id="resetBox">|<\\/span>$/g, ""); wireReset(); };
  };
}
function next(){ state.index++; save(); render(); window.scrollTo(0, 0); }
function zoom(src){
  const o = document.createElement("div"); o.className = "lightbox";
  o.innerHTML = '<img src="' + src + '" alt=""><span>' + esc(T.zoom_close) + '</span>'; o.onclick = () => o.remove(); document.body.appendChild(o);
}
function mailto(){
  const fill = s => s.replace("{version}", D.version).replace("{platform}", D.platform).replace("{file}", D.resultFile);
  return "mailto:" + D.email + "?subject=" + encodeURIComponent(fill(T.mail_subject)) + "&body=" + encodeURIComponent(fill(T.mail_body));
}
function plain(h){ const d = document.createElement("div"); d.innerHTML = h; return d.textContent; }
function resultText(){
  const lines = [D.title, T.result_head + " · " + new Date().toLocaleString(), T.time_total.replace("{min}", mins(totalMs())), ""];
  if (state.mode) lines.push(T.mode_label + ": " + T.modes[state.mode][0]);
  D.fields.forEach((f, i) => { const v = (state.device[i] || "").trim(); if (v) lines.push(f.label + ": " + v); });
  lines.push("");
  D.tests.forEach(t => {
    if (!TESTS.includes(t)) { lines.push(t.id + " " + plain(t.title) + ": " + T.status.na); return; }
    const r = state.results[t.id];
    lines.push(t.id + " " + plain(t.title) + ": " + T.status[r ? r.s : "open"] + (r && r.text ? " – " + r.text : "") + (r ? " (" + oneTime(state.secs[t.id] || 0) + ")" : ""));
  });
  return lines.join("\\r\\n") + "\\r\\n";
}
function renderSummary(){
  const c = counts();
  view.innerHTML = '<section class="card"><div class="hero">' + mascot("thumbs") + '<div><div class="eyebrow">' + esc(T.done_eyebrow) + '</div><h1>' + esc(T.done) + '</h1><p>' + esc(T.done_sub) + '</p></div></div>'
    + '<p>' + esc(T.time_total.replace("{min}", mins(totalMs()))) + '</p>'
    + '<div class="sum"><div class="s-ok"><b>' + c.ok + '</b>' + esc(T.labels.ok) + '</div><div class="s-bad"><b>' + c.bad + '</b>' + esc(T.labels.bad) + '</div><div class="s-skip"><b>' + c.skip + '</b>' + esc(T.labels.skip) + '</div></div>'
    + '<div style="overflow-x:auto"><table><tbody>' + TESTS.map(t => { const r = state.results[t.id], k = r ? r.s : "open";
        return '<tr><td>' + esc(t.id) + '</td><td>' + t.title + (r && r.text ? '<br><span style="color:var(--muted)">' + esc(r.text) + '</span>' : '') + '</td><td><span class="tag t-' + k + '">' + esc(T.labels[k]) + '</span></td></tr>'; }).join("") + '</tbody></table></div>'
    + '<p class="q">' + esc(T.device) + '</p><div class="fields">' + D.fields.map((f, i) => '<label>' + esc(f.label)
        + (f.choices.length ? '<select data-f="' + i + '"><option value=""></option>' + f.choices.map(x => '<option' + (state.device[i] === x ? ' selected' : '') + '>' + esc(x) + '</option>').join("") + '</select>'
                            : '<input data-f="' + i + '" placeholder="' + esc(f.hint) + '" value="' + esc(state.device[i] || "") + '">') + '</label>').join("") + '</div>'
    + '<p>' + esc(T.final) + '</p><button class="primary" id="saveRes">' + esc(T.save) + '</button><p class="saved" id="saved" hidden></p>'
    + '<a class="mailbtn" href="' + esc(mailto()) + '">✉ ' + esc(T.mail_btn) + '<span>' + esc(D.email) + '</span></a><div class="tip" style="margin-top:0">' + esc(T.mail_hint) + '</div>'
    + '<div class="nav"><button class="linkbtn" id="back">' + esc(T.back_last) + '</button>' + resetBox() + '</div></section>';
  view.querySelectorAll("[data-f]").forEach(el => el.oninput = el.onchange = () => { state.device[el.dataset.f] = el.value; save(); });
  document.getElementById("saveRes").onclick = () => {
    const blob = new Blob(["﻿" + resultText()], { type:"text/plain;charset=utf-8" }), a = document.createElement("a");
    a.href = URL.createObjectURL(blob); a.download = D.resultFile; document.body.appendChild(a); a.click(); a.remove();
    setTimeout(() => URL.revokeObjectURL(a.href), 1000);
    const s = document.getElementById("saved"); s.textContent = T.saved.replace("{file}", D.resultFile); s.hidden = false;
  };
  document.getElementById("back").onclick = () => { state.index = TESTS.length - 1; save(); render(); };
  wireReset();
}

document.addEventListener("keydown", e => {
  if (e.key === "Escape") { const o = document.querySelector(".lightbox"); if (o) { o.remove(); return; } }
  if (document.querySelector(".lightbox")) return;
  if (!state.started || !state.prepDone || state.index >= TESTS.length) return;
  if (["TEXTAREA","SELECT","INPUT"].includes(document.activeElement && document.activeElement.tagName)) return;
  const m = { "1":"ok", "2":"bad", "3":"skip" }[e.key]; if (m) choose(TESTS[state.index], m);
});
render();
</script></body></html>
"""


CHROMIUM_CANDIDATES = ["chromium", "google-chrome", "google-chrome-stable", "/usr/bin/chromium"]


def render_pdf(html_path, pdf_path):
    chromium = next((c for c in CHROMIUM_CANDIDATES if shutil.which(c) or pathlib.Path(c).is_file()), None)
    if chromium is None:
        raise ValueError("Kein Chromium/Chrome für das PDF gefunden")
    subprocess.run(
        [chromium, "--headless", "--no-sandbox", "--no-pdf-header-footer",
         f"--print-to-pdf={pdf_path}", html_path.resolve().as_uri()],
        check=True, capture_output=True,
    )


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    flags = [a for a in argv if a.startswith("--")]
    lang = next((f.split("=", 1)[1] for f in flags if f.startswith("--lang=")), "de")
    if len(args) < 3:
        print("usage: testanleitung.py <version> <windows|macos|linux> <out_dir> [--pdf] [--lang=de|en] [--assistant-only]",
              file=sys.stderr)
        return 1
    version, platform, out_dir = args[0], args[1], args[2]
    try:
        unknown = [f for f in flags if f not in ("--pdf", "--assistant-only") and not f.startswith("--lang=")]
        if unknown:
            raise ValueError(f"Unbekannte Option: {' '.join(unknown)}")
        if platform not in PLATFORM_NAMES:
            raise ValueError(f"Unbekannte Plattform: {platform}")
        if lang not in LANGS:
            raise ValueError(f"Unbekannte Sprache: {lang}")
        build(version, platform, out_dir, pdf="--pdf" in flags, lang=lang,
              assistant_only="--assistant-only" in flags)
    except ValueError as e:
        print(f"testanleitung: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
