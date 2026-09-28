import base64, json, pathlib, tempfile, unittest
from openpyxl import load_workbook
import testanleitung as ta

SAMPLE = """# 3 · Einrichten

## E2 · Erster Start
Nur: windows
Schritte:
1. Die App starten.
2. Den Ordner wählen.
Erwartet: Leerer Katalog.
Hinweis: Nur beim ersten Mal.

## E3 · Backup
Schritte:
1. Backup einspielen.
Erwartet: Filament ist da.
"""

class Parse(unittest.TestCase):
    def test_chapters_and_scenarios(self):
        ch = ta.parse_scenarios(SAMPLE)
        self.assertEqual(ch[0].title, "3 · Einrichten")
        e2 = ch[0].scenarios[0]
        self.assertEqual((e2.id, e2.title, e2.steps, e2.expected, e2.note, e2.only),
                         ("E2", "Erster Start", ["Die App starten.", "Den Ordner wählen."], "Leerer Katalog.", "Nur beim ersten Mal.", {"windows"}))
        self.assertEqual(ch[0].scenarios[1].only, set())

    def test_missing_expectation_names_the_id(self):
        with self.assertRaisesRegex(ValueError, "E9"):
            ta.parse_scenarios("# K\n\n## E9 · X\nSchritte:\n1. a\n")

class ParseErrors(unittest.TestCase):
    def test_scenario_without_chapter(self):
        with self.assertRaisesRegex(ValueError, "E1"):
            ta.parse_scenarios("## E1 · X\nSchritte:\n1. a\nErwartet: b\n")

    def test_text_outside_a_scenario_is_an_error(self):
        with self.assertRaisesRegex(ValueError, "Einleitung"):
            ta.parse_scenarios("# K\n\nEinleitung zum Kapitel.\n\n## E1 · X\nSchritte:\n1. a\nErwartet: b\n")

    def test_last_scenario_of_a_chapter_is_checked(self):
        with self.assertRaisesRegex(ValueError, "E1"):
            ta.parse_scenarios("# K\n\n## E1 · X\nSchritte:\n1. a\n\n# L\n\n## E2 · Y\nSchritte:\n1. a\nErwartet: b\n")


class Fill(unittest.TestCase):
    def test_unknown_placeholder_fails(self):
        with self.assertRaisesRegex(ValueError, "unbekannt"):
            ta.fill("{{version}} {{unbekannt}}", {"version": "1"})

    def test_platform_values(self):
        v = ta.platform_values("0.15.0-2", "windows")
        self.assertEqual(v["paket"], "3MF-Katalog-Manager-Preview-0.15.0-2-Windows-x64.msi")
        self.assertIn("com.thebexxs.mfkatalogmanager.preview", v["datenordner"])
        self.assertIn("version=0.15.0-2", v["formular"])
        # Log paths match Tauri's app_log_dir (see tauri-plugin-log's
        # TargetKind::LogDir), which is NOT "<data folder>/logs" on Windows/macOS.
        self.assertEqual(v["logordner"], r"%LOCALAPPDATA%\com.thebexxs.mfkatalogmanager.preview\logs")
        macos = ta.platform_values("0.15.0-2", "macos")
        self.assertEqual(macos["logordner"], "~/Library/Logs/com.thebexxs.mfkatalogmanager.preview")
        linux = ta.platform_values("0.15.0-2", "linux")
        self.assertEqual(linux["logordner"], "~/.local/share/com.thebexxs.mfkatalogmanager.preview/logs")

class FillScenarioText(unittest.TestCase):
    def test_substitutes_placeholders_in_place(self):
        ch = ta.parse_scenarios("# K\n\n## U5 · Sicherung\nSchritte:\n1. Öffne {{datenordner}}.\nErwartet: Datei {{version}}.db liegt dort.\n")
        ta.fill_scenario_text(ch, {"datenordner": "/x/y", "version": "0.15.0-2"})
        s = ch[0].scenarios[0]
        self.assertEqual(s.steps, ["Öffne /x/y."])
        self.assertEqual(s.expected, "Datei 0.15.0-2.db liegt dort.")

    def test_unknown_placeholder_in_scenario_text_fails(self):
        ch = ta.parse_scenarios("# K\n\n## U5 · X\nSchritte:\n1. a\nErwartet: {{unbekannt}}\n")
        with self.assertRaisesRegex(ValueError, "unbekannt"):
            ta.fill_scenario_text(ch, {"version": "1"})

class Changelog(unittest.TestCase):
    TEXT = "# Changelog\n\n## [Unreleased]\n\n### Added\n- **Update:** EN text\n\n## [0.14.0]\n- old\n\n# Changelog (Deutsch)\n\n## [Unreleased]\n\n### Hinzugefügt\n- **Update:** DE-Text\n\n## [0.14.0]\n- alt\n"

    def test_german_part_of_unreleased(self):
        de = ta.changelog_unreleased(self.TEXT, "de")
        self.assertIn("DE-Text", de)
        self.assertNotIn("alt", de)
        self.assertNotIn("EN text", de)

    def test_english_part_of_unreleased(self):
        en = ta.changelog_unreleased(self.TEXT, "en")
        self.assertIn("EN text", en)
        self.assertNotIn("old", en)
        self.assertNotIn("DE-Text", en)

class ScenarioHtml(unittest.TestCase):
    def test_backticks_become_code_and_no_raw_html_passes(self):
        s = ta.Scenario(
            id="P1", title="Mit `code`",
            steps=["Nutze `$f = [System.IO.File]::Open(\"<Katalog>\\x.3mf\")`."],
            expected="Zeigt `<script>alert(1)</script>` nicht aus.",
            note="Auch `hier`.", only=set(),
        )
        out = ta._scenario_html(s)
        self.assertIn("<code>code</code>", out)
        self.assertIn("<code>$f = [System.IO.File]::Open(&quot;&lt;Katalog&gt;\\x.3mf&quot;)</code>", out)
        self.assertIn("<code>hier</code>", out)
        # The angle brackets inside a backtick span must stay escaped text, not
        # turn into real tags - that was the whole bug (raw HTML from user text).
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", out)
        self.assertNotIn("<script>", out)
        self.assertNotIn("<Katalog>", out)

SAMPLE_EN = """# 3 · Setup

## E2 · First start
Only: windows
Steps:
1. Start the app.
2. Choose the folder.
Expected: Empty catalog.
Note: Only the first time.

## E3 · Backup
Steps:
1. Restore the backup.
Expected: Filament is there.
"""


def make_root(t, english=SAMPLE_EN):
    root = pathlib.Path(t)
    for folder in ("docs/tests/bausteine", "docs/tests/bausteine/en"):
        (root / folder).mkdir(parents=True, exist_ok=True)
        for platform in ("windows", "linux"):
            for name in (f"installieren-{platform}", "daten", "fehler-melden", f"deinstallieren-{platform}"):
                (root / folder / f"{name}.md").write_text(f"Baustein {name} {{{{version}}}} {{{{formular}}}}\n")
    (root / "docs/tests/0.15.0-2.md").write_text(SAMPLE)
    if english is not None:
        (root / "docs/tests/0.15.0-2.en.md").write_text(english)
    (root / "CHANGELOG.md").write_text("## [Unreleased]\n### Added\n- New\n\n# Changelog (Deutsch)\n\n## [Unreleased]\n### Hinzugefügt\n- Neu\n")
    return root


def assistant_data(path):
    text = path.read_text(encoding="utf-8")
    start = text.index('<script id="data" type="application/json">') + len('<script id="data" type="application/json">')
    return text, json.loads(text[start:text.index("</script>", start)])


class Images(unittest.TestCase):
    WITH_IMAGES = """# 3 · Einrichten

## E2 · Erster Start
Schritte:
1. Zahnrad klicken.
   Bild: zahnrad · ① Zahnrad
2. Weiter.
Erwartet: Leerer Katalog.
Bild-Ergebnis: leer · so sieht es aus
"""

    def test_parse_step_and_result_images(self):
        s = ta.parse_scenarios(self.WITH_IMAGES)[0].scenarios[0]
        self.assertEqual(s.steps, ["Zahnrad klicken.", "Weiter."])
        self.assertEqual(s.step_images, {0: ("zahnrad", "① Zahnrad")})
        self.assertEqual(s.result_image, ("leer", "so sieht es aus"))

    def test_platform_picture_wins_and_both_outputs_embed_it(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            (root / "docs/tests/0.15.0-2.md").write_text(self.WITH_IMAGES)
            (root / "docs/tests/0.15.0-2.en.md").write_text(self.WITH_IMAGES.replace("Schritte:", "Steps:").replace("Erwartet:", "Expected:")
                                                            .replace("   Bild:", "   Image:").replace("Bild-Ergebnis:", "Result image:"))
            bilder = root / "docs/tests/bilder"
            bilder.mkdir()
            (bilder / "zahnrad.webp").write_bytes(b"ALL")
            (bilder / "zahnrad.windows.webp").write_bytes(b"WIN")
            (bilder / "leer.webp").write_bytes(b"EMPTY")
            html_path, _ = ta.build("0.15.0-2", "windows", root / "out", root=root)
            win = "data:image/webp;base64," + base64.b64encode(b"WIN").decode()
            self.assertIn(win, html_path.read_text())
            self.assertIn("So sieht es richtig aus", html_path.read_text())
            _, data = assistant_data(root / "out" / "Testassistent-0.15.0-2-Windows.html")
            self.assertEqual(data["tests"][0]["steps"][0]["img"], win)
            self.assertEqual(data["tests"][0]["steps"][1]["img"], "")
            self.assertEqual(data["tests"][0]["resultImg"], "data:image/webp;base64," + base64.b64encode(b"EMPTY").decode())
            ta.build("0.15.0-2", "linux", root / "out", root=root)
            _, linux = assistant_data(root / "out" / "Testassistent-0.15.0-2-Linux.html")
            self.assertEqual(linux["tests"][0]["steps"][0]["img"], "data:image/webp;base64," + base64.b64encode(b"ALL").decode())

    def test_missing_picture_fails_with_the_scenario_id(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            (root / "docs/tests/0.15.0-2.md").write_text(self.WITH_IMAGES)
            with self.assertRaisesRegex(ValueError, "E2.*Bild fehlt"):
                ta.build("0.15.0-2", "windows", root / "out", root=root)


class Assistant(unittest.TestCase):
    def test_german_assistant_is_written_next_to_the_guide(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            ta.build("0.15.0-2", "windows", root / "out", root=root)
            text, data = assistant_data(root / "out" / "Testassistent-0.15.0-2-Windows.html")
            self.assertEqual([s["id"] for s in data["tests"]], ["E2", "E3"])
            self.assertEqual([s["t"] for s in data["tests"][0]["steps"]], ["Die App starten.", "Den Ordner wählen."])
            self.assertEqual(data["tests"][0]["note"], "Nur beim ersten Mal.")
            self.assertEqual(data["resultFile"], "Testergebnis-0.15.0-2-Windows.txt")
            self.assertIn("Baustein installieren-windows 0.15.0-2", data["install"])
            self.assertIn("Virenschutz", [f["label"] for f in data["fields"]])
            self.assertIn('<html lang="de">', text)

    def test_english_assistant(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            ta.build("0.15.0-2", "windows", root / "out", root=root, lang="en")
            text, data = assistant_data(root / "out" / "Test-Assistant-0.15.0-2-Windows.html")
            self.assertEqual(data["tests"][0]["title"], "First start")
            self.assertEqual(data["text"]["ok"], "Works")
            self.assertEqual(data["resultFile"], "Test-Result-0.15.0-2-Windows.txt")

    def test_assistant_lists_only_the_platforms_scenarios(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            ta.build("0.15.0-2", "linux", root / "out", root=root)
            _, data = assistant_data(root / "out" / "Testassistent-0.15.0-2-Linux.html")
            self.assertEqual([s["id"] for s in data["tests"]], ["E3"])
            self.assertIn("Sitzung", [f["label"] for f in data["fields"]])

    def test_scenario_text_cannot_end_the_script_or_inject_html(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            (root / "docs/tests/0.15.0-2.md").write_text(SAMPLE.replace("Leerer Katalog.", "Kein `</script><b>x</b>` zu sehen."))
            ta.build("0.15.0-2", "windows", root / "out", root=root)
            text, data = assistant_data(root / "out" / "Testassistent-0.15.0-2-Windows.html")
            self.assertEqual(text.count("</script>"), 2)
            self.assertIn("<code>&lt;/script&gt;&lt;b&gt;x&lt;/b&gt;</code>", data["tests"][0]["expected"])


def sheet_rows(path):
    return [[c for c in row] for row in load_workbook(path).active.iter_rows(values_only=True)]


class Build(unittest.TestCase):
    def test_german_guide_points_to_the_sheet(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            out = root / "out"
            html_path, sheet_path = ta.build("0.15.0-2", "windows", out, root=root)
            html = html_path.read_text()
            self.assertEqual(html_path.name, "Testanleitung-0.15.0-2-Windows.html")
            self.assertEqual(sheet_path.name, "Ergebnisbogen-0.15.0-2-Windows.xlsx")
            self.assertIn("Erster Start", html)
            self.assertIn("Baustein daten 0.15.0-2 https://3mfkatalog.de/fehler-melden.html", html)
            self.assertIn("Ergebnisbogen-0.15.0-2-Windows.xlsx", html)
            self.assertIn("Zeile E2", html)
            self.assertNotIn("&#9744;", html)  # no checkboxes that can't be ticked in a PDF
            self.assertIn("Neu", html)

    def test_english_guide_and_sheet(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            html_path, sheet_path = ta.build("0.15.0-2", "windows", root / "out", root=root, lang="en")
            html = html_path.read_text()
            self.assertEqual(html_path.name, "Test-Guide-0.15.0-2-Windows.html")
            self.assertIn("First start", html)
            self.assertIn("Steps:", html)
            self.assertIn("https://3mfkatalog.de/en/report-a-bug.html?version=0.15.0-2&amp;os=windows", html)
            self.assertIn("New", html)
            self.assertNotIn("Neu", html)
            rows = sheet_rows(sheet_path)
            self.assertIn(("No.", "Test", "Result", "Note"), [tuple(r) for r in rows])

    def test_sheet_lists_only_the_platforms_scenarios(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            _, windows = ta.build("0.15.0-2", "windows", root / "out", root=root)
            _, linux = ta.build("0.15.0-2", "linux", root / "out", root=root)
            ids = lambda path: [r[0] for r in sheet_rows(path) if r[0] in ("E2", "E3")]
            self.assertEqual(ids(windows), ["E2", "E3"])
            self.assertEqual(ids(linux), ["E3"])
            labels = [r[0] for r in sheet_rows(linux)]
            self.assertIn("Sitzung", labels)
            self.assertNotIn("Windows-Version", labels)

    def test_result_list_and_summary(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t)
            _, sheet_path = ta.build("0.15.0-2", "windows", root / "out", root=root)
            ws = load_workbook(sheet_path).active
            lists = [dv.formula1 for dv in ws.data_validations.dataValidation]
            self.assertIn('"OK,Fehler,Übersprungen"', lists)
            formulas = [ws.cell(r, 3).value for r in range(1, ws.max_row + 1)]
            self.assertTrue(any(str(f).startswith("=COUNTIF(") for f in formulas))

    def test_english_scenarios_must_match_the_german_ids(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t, english=SAMPLE_EN.replace("## E3 · Backup", "## E4 · Backup"))
            with self.assertRaisesRegex(ValueError, "passen nicht"):
                ta.build("0.15.0-2", "windows", root / "out", root=root, lang="en")

    def test_missing_english_scenario_file(self):
        with tempfile.TemporaryDirectory() as t:
            root = make_root(t, english=None)
            with self.assertRaisesRegex(ValueError, "Szenario-Datei fehlt"):
                ta.build("0.15.0-2", "windows", root / "out", root=root, lang="en")


class EnglishParse(unittest.TestCase):
    def test_english_keywords(self):
        ch = ta.parse_scenarios(SAMPLE_EN)
        self.assertEqual(ch[0].scenarios[0].only, {"windows"})
        self.assertEqual(ch[0].scenarios[0].steps, ["Start the app.", "Choose the folder."])
        self.assertEqual(ch[0].scenarios[0].note, "Only the first time.")


if __name__ == "__main__":
    unittest.main()
